# -*- coding: utf-8 -*-
"""멀티에셋 투자 유니버스 — universe/universe.json(원본 데이터)을 읽고 엑셀로 내보낸다.

universe/universe.json : 예시 문서(자산·파생상품·결합 전략)를 옮긴 원본. 고칠 때는 이 파일을 고친다.
universe/equity.json   : 주식 항목별 데이터 연결(티커)과 전략 준비 상태 — 주식 대시보드(equity.py)가 쓴다.

사용법:
  python universe.py            # universe/멀티에셋_유니버스_전략.xlsx 다시 만들기
"""
import json
import sys
from pathlib import Path

import config

DIR = config.BASE_DIR / "universe"
SRC = DIR / "universe.json"
EQ = DIR / "equity.json"
XLSX = DIR / "멀티에셋_유니버스_전략.xlsx"


def load() -> dict:
    return json.loads(SRC.read_text(encoding="utf-8"))


def load_equity() -> dict:
    return json.loads(EQ.read_text(encoding="utf-8"))


def equity_view(u: dict | None = None, eq: dict | None = None) -> dict:
    """주식 부분만: 기초자산 항목(+ 연결된 데이터), 주식 관련 파생상품, 결합 전략(+ 준비 상태)."""
    u = u or load()
    eq = eq or load_equity()
    series = eq["series"]
    items = []
    for it in (x for x in u["plain_vanilla"] if x["asset"] == "주식"):
        items.append({**it, "series": [{"key": s["key"], "name": s["name"], "ticker": s["ticker"]} for s in series if s.get("item") == it["id"]]})
    strats = []
    for s in (x for x in u["strategies"] if x["asset"] == "주식"):
        r = eq["strategies"].get(s["id"], {})
        strats.append({**s, "ready": r.get("ready"), "ready_text": eq["readiness_levels"].get(r.get("ready"), ""),
                       "data": r.get("data", []), "backtest": r.get("backtest", "")})
    return {"items": items, "derivatives": [d for d in u["derivatives"] if d.get("equity")],
            "strategies": strats, "missing": eq.get("missing", []),
            "notes": [n for n in u["meta"].get("notes", []) if "주식" in n["where"] or n["where"] == "개요"],
            "meta": {k: u["meta"][k] for k in ("title", "concept", "goal", "horizon")}}


def export_excel(path: Path = XLSX) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    u, eq = load(), load_equity()
    F = "Arial"
    head, fill = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="3A3A3C")
    base, bold, sub = Font(name=F, size=10), Font(name=F, size=10, bold=True), Font(name=F, size=9, color="666666")
    wb = Workbook()

    def table(ws, row, cols, rows, widths):
        for c, h in enumerate(cols, 1):
            cell = ws.cell(row, c, h); cell.font, cell.fill = head, fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, r in enumerate(rows, start=row + 1):
            for c, v in enumerate(r, 1):
                cell = ws.cell(i, c, v); cell.font = base
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + c)].width = w
        ws.freeze_panes = ws.cell(row + 1, 1)

    ws = wb.active; ws.title = "개요"
    m = u["meta"]
    ws["A1"] = m["title"]; ws["A1"].font = Font(name=F, size=14, bold=True)
    for i, (k, v) in enumerate([("구성", m["concept"]), ("목표", m["goal"]), ("투자시계", m["horizon"]),
                                ("기초자산", " · ".join(a["name"] + (f"({a['desc']})" if a.get("desc") else "") for a in u["assets"])),
                                ("파생상품(결합 수단)", " · ".join(u["derivative_types"])),
                                ("규모", f"기초자산 {len(u['plain_vanilla'])}개 항목 · 파생상품 {len(u['derivatives'])}개 · 결합 전략 {len(u['strategies'])}개"),
                                ("출처", m["source"])], start=3):
        ws.cell(i, 1, k).font = bold; ws.cell(i, 2, v).font = base
    ws.cell(11, 1, "메모").font = bold
    for i, n in enumerate(m.get("notes", []), start=12):
        ws.cell(i, 1, n["where"]).font = base; ws.cell(i, 2, f"{n['text']} ({n['kind']})").font = base
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 110

    ws = wb.create_sheet("기초자산")
    table(ws, 1, ["ID", "자산군", "세부 항목", "예시"], [(x["id"], x["asset"], x["item"], x["examples"]) for x in u["plain_vanilla"]], [11, 9, 18, 90])
    ws = wb.create_sheet("파생상품")
    table(ws, 1, ["ID", "파생 분류", "세부 항목", "예시", "주식 관련"],
          [(x["id"], x["type"], x["item"], x["examples"], "○" if x.get("equity") else "") for x in u["derivatives"]], [11, 18, 20, 80, 9])
    ws = wb.create_sheet("결합전략")
    table(ws, 1, ["ID", "자산군", "파생 분류", "전략", "구조", "수익을 내는 시나리오", "손익 형태", "주요 리스크"],
          [(x["id"], x["asset"], x["type"], x["name"], x["structure"], x["scenario"], x["payoff"], x["risk"]) for x in u["strategies"]],
          [11, 8, 18, 20, 32, 30, 12, 26])

    v = equity_view(u, eq)
    ws = wb.create_sheet("주식")
    ws["A1"] = "주식 — 유니버스·데이터·전략 준비 상태"; ws["A1"].font = Font(name=F, size=14, bold=True)
    ws["A2"] = eq["meta"]["note"]; ws["A2"].font = sub
    table(ws, 4, ["ID", "세부 항목", "예시", "대시보드 데이터(티커)"],
          [(x["id"], x["item"], x["examples"], ", ".join(f"{s['name']} [{s['ticker']}]" for s in x["series"]) or "-") for x in v["items"]],
          [11, 20, 34, 90])
    r0 = 6 + len(v["items"]) + 1
    ws.cell(r0, 1, "결합 전략 (주식)").font = bold
    rows = [(s["id"], s["type"], s["name"], s["structure"], s["scenario"], s["payoff"], s["risk"],
             eq["readiness_levels"].get(s["ready"], ""), s["backtest"]) for s in v["strategies"]]
    cols = ["ID", "파생 분류", "전략", "구조", "수익을 내는 시나리오", "손익 형태", "주요 리스크", "백테스트 준비", "계산 방법"]
    for c, h in enumerate(cols, 1):
        cell = ws.cell(r0 + 1, c, h); cell.font, cell.fill = head, fill
    for i, r in enumerate(rows, start=r0 + 2):
        for c, val in enumerate(r, 1):
            cell = ws.cell(i, c, val); cell.font = base; cell.alignment = Alignment(vertical="top", wrap_text=True)
    for c, w in zip("EFGHI", (30, 12, 24, 22, 60)):
        ws.column_dimensions[c].width = w
    r1 = r0 + 3 + len(rows)
    ws.cell(r1, 1, "아직 연결 안 된 데이터").font = bold
    for i, x in enumerate(v["missing"], start=r1 + 1):
        ws.cell(i, 2, x["name"]).font = base; ws.cell(i, 3, x["need"]).font = base; ws.cell(i, 4, x["source"]).font = base
    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    u = load()
    print(f"기초자산 {len(u['plain_vanilla'])} · 파생상품 {len(u['derivatives'])} · 결합 전략 {len(u['strategies'])}")
    print("엑셀:", export_excel())
