# -*- coding: utf-8 -*-
"""보유 엑셀(시험 포트) → 시가 평가 → data/positions.json

엑셀의 입력값(수량·진입가·진입환율·승수·방향·베타·듀레이션·헤지·목표·손절·엔진구분)만 읽고,
금액·손익은 여기서 다시 계산한다(엑셀 수식의 캐시값에 의존하지 않음).

가격:
- 출처 Refinitiv(refinitiv_src 가 아래 표기를 RIC 으로 바꾼다): KRX 주식·ETF <코드>.KS / 미국 ETF 티커 / 환율 KRW=X, JPYKRW=X
- KOSPI200 선물 → KOSPI200 지수 현물(^KS200), 달러선물 → USD/KRW 현물, 엔선물 → 엔/원 현물×100 (베이시스 무시)
- 10년 국채선물 → 최신 국고10y 로 할인한 표면 5% 10년 표준물 이론가
손익 분해: 자산 손익(가격 변화) / 환 손익(외화 자산의 미헤지분 + 환 오버레이). 헤지 비용·거래비용은 미반영.

사용법:
  python book.py                       # 기본 파일(data/시험포트.xlsx) → data/positions.json
  python book.py 다른파일.xlsx --limits  # 엑셀 '한도' 시트로 data/limits.json 도 갱신(체험용)
  python book.py --dry-run             # 쓰지 않고 평가 결과만 출력
"""
import argparse
import datetime
import json
import os
import sys
import warnings
from pathlib import Path

import config

DEFAULT_BOOK = config.DATA_DIR / "시험포트.xlsx"
KRX_MARKETS = ("KRX 유가증권", "KRX ETF")


# ── 엑셀 읽기 ───────────────────────────────────────────────────
def read_book(path: Path) -> tuple[list[dict], dict[str, dict], dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)   # 입력 칸만 읽으므로 수식 캐시가 없어도 된다
    legs = []
    ws = wb["구성종목"]
    for r in range(5, ws.max_row + 1):
        rid = ws[f"A{r}"].value
        if not rid:
            continue
        fx_cell = ws[f"J{r}"].value
        legs.append({
            "row": rid, "pos": ws[f"B{r}"].value, "name": ws[f"C{r}"].value, "code": str(ws[f"D{r}"].value),
            "market": ws[f"E{r}"].value, "ccy": ws[f"F{r}"].value, "side": ws[f"G{r}"].value,
            "entry_px": float(ws[f"I{r}"].value), "entry_fx": fx_cell, "mult": float(ws[f"K{r}"].value),
            "qty": float(ws[f"L{r}"].value), "beta": ws[f"P{r}"].value, "dur": ws[f"Q{r}"].value,
        })
    pos = {}
    ws = wb["포지션"]
    for r in range(5, ws.max_row + 1):
        pid = ws[f"A{r}"].value
        if not pid or pid == "합계" or not ws[f"R{r}"].value:
            continue
        pos[pid] = {
            "id": pid, "asset_ko": ws[f"B{r}"].value, "strategy": ws[f"C{r}"].value, "how": ws[f"D{r}"].value,
            "name": ws[f"E{r}"].value, "hedge": ws[f"I{r}"].value, "target": ws[f"K{r}"].value, "stop": ws[f"L{r}"].value,
            "engine": ws[f"R{r}"].value, "ccy": ws[f"S{r}"].value, "sdur": ws[f"T{r}"].value,
            "sbeta": ws[f"U{r}"].value, "eq_pct": ws[f"V{r}"].value,
            "thesis": ws[f"O{r}"].value, "invalidation": ws[f"P{r}"].value, "evidence": ws[f"Q{r}"].value,
        }
    # 가정 시트: 진입 환율(외화 다리의 J 칸이 이 셀을 참조하는 수식이라 값으로 다시 읽는다), 달러선물 목표·손절 환율
    ga = {}
    ws = wb["가정·출처"]
    for r in range(5, ws.max_row + 1):
        k, v = ws[f"A{r}"].value, ws[f"B{r}"].value
        if k and isinstance(v, (int, float)):
            ga[k] = float(v)
    lim = {}
    if "한도(체험용 가안)" in wb.sheetnames:
        ws = wb["한도(체험용 가안)"]
        for r in range(5, 9):
            k, v = ws[f"A{r}"].value, ws[f"B{r}"].value
            if k and isinstance(v, (int, float)):
                lim[k] = float(v)
    ga["_limits"] = lim
    for l in legs:
        if not isinstance(l["entry_fx"], (int, float)):
            l["entry_fx"] = ga["USD/KRW 진입 환율"] if l["ccy"] == "USD" else 1.0
        l["entry_fx"] = float(l["entry_fx"])
    return legs, pos, ga


# ── 가격 ────────────────────────────────────────────────────────
def _ktbf_price(y_pct: float) -> float:
    c, r = 2.5, y_pct / 200
    return sum((c + (100 if t == 20 else 0)) / (1 + r) ** t for t in range(1, 21))


def price_source(l: dict) -> str:
    n = l["name"]
    if "KOSPI200 선물" in n:
        return "^KS200"
    if "국채선물" in n:
        return "KTB10"
    if "달러선물" in n:
        return "KRW=X"
    if "엔선물" in n:
        return "JPYKRW=X"
    if l["market"] in KRX_MARKETS:
        return f"{l['code']}.KS"
    return l["code"]


def fetch_prices(symbols: set[str]) -> dict[str, tuple[float, str]]:
    warnings.filterwarnings("ignore")
    import refinitiv_src

    out = {}
    for s in sorted(symbols):
        if s == "KTB10":
            import analytics
            v = analytics.latest("ktb_10y")
            if v:
                out[s] = (_ktbf_price(v[0]), v[1])
            continue
        try:
            if s in refinitiv_src.YF_ONLY:
                import yfinance as yf
                h = yf.Ticker(s).history(period="7d")
                h = h[h["Close"].notna()]
                if len(h):
                    out[s] = (float(h["Close"].iloc[-1]), h.index[-1].date().isoformat())
            else:
                v, d = refinitiv_src.last_close(s)
                out[s] = (v, d.isoformat())
        except Exception as e:  # noqa: BLE001 — 한 종목 실패는 진입가 유지로 처리
            print(f"  경고: {s} 가격 조회 실패 — {e}")
    return out


# ── 평가 ────────────────────────────────────────────────────────
def mark(legs: list[dict], pos: dict, prices: dict, ga: dict) -> tuple[list[dict], list[dict]]:
    fx_now = prices.get("KRW=X", (ga["USD/KRW 진입 환율"], None))[0]
    rows = []
    for l in legs:
        src = price_source(l)
        p_now, p_date = prices.get(src, (l["entry_px"], None))
        sign = -1 if l["side"] == "숏" else 1
        fxn = fx_now if l["ccy"] == "USD" else 1.0
        if src == "JPYKRW=X":
            p_now *= 100                           # 엔선물은 100엔당 원
        p = pos[l["pos"]]
        h = float(p["hedge"] or 0) if l["ccy"] == "USD" else 0.0
        entry_val = l["qty"] * l["entry_px"] * l["mult"] * l["entry_fx"] / 1e8
        now_val = l["qty"] * p_now * l["mult"] * fxn / 1e8
        if p["engine"] == "fx":                    # 오버레이: 가격 자체가 환율 → 전부 환 손익
            pnl_asset, pnl_fx = 0.0, sign * l["qty"] * (p_now - l["entry_px"]) * l["mult"] / 1e8
        else:
            pnl_asset = sign * l["qty"] * (p_now - l["entry_px"]) * l["mult"] * fxn / 1e8
            fx_part = sign * l["qty"] * l["entry_px"] * l["mult"] * (fxn - l["entry_fx"]) / 1e8
            pnl_fx = (1 - h) * fx_part             # 헤지분 환손익은 선물환이 상쇄
        rows.append({**l, "src": src, "px_now": round(p_now, 4), "px_date": p_date or "진입가 유지",
                     "sign": sign, "entry_bn": round(entry_val, 2), "now_bn": round(now_val, 2),
                     "pnl_asset_bn": round(pnl_asset, 2), "pnl_fx_bn": round(pnl_fx, 2)})

    today = datetime.date.today().isoformat()
    out = []
    for pid, p in pos.items():
        lr = [r for r in rows if r["pos"] == pid]
        if not lr:
            continue
        longs = [r for r in lr if r["sign"] > 0]
        base = sum(r["entry_bn"] for r in longs)
        mv = sum(r["now_bn"] for r in longs)
        pnl = sum(r["pnl_asset_bn"] for r in lr)
        pnl_fx = sum(r["pnl_fx_bn"] for r in lr)
        dates = [r["px_date"] for r in lr if r["px_date"] != "진입가 유지"]
        net_beta = sum(r["sign"] * r["now_bn"] * float(r["beta"] or 0) for r in lr) / mv if mv else 0.0
        net_dur = sum(r["sign"] * r["now_bn"] * float(r["dur"] or 0) for r in lr) / mv if mv else 0.0
        stop_pct = round(float(p["stop"]) * 100, 2) if p["stop"] is not None else None
        if pid == "F1" and stop_pct is None:       # 달러선물: 엑셀 수식 칸 → 가정 시트 환율로 다시 계산
            e = ga["USD/KRW 진입 환율"]
            stop_pct = round((ga["달러선물 손절 환율"] / e - 1) * 100, 2)
            p["target"] = ga["달러선물 목표 환율"] / e - 1
        rec = {
            "id": pid, "name": p["name"], "short_name": p["name"].split(" (")[0][:24], "type": "bucket",
            "asset": p["engine"], "group": p["asset_ko"], "sub": p["strategy"], "strategy": p["strategy"], "management": p["how"],
            "pnl_as_of": min(dates) if dates else today, "rate_pnl_in_scope": True,
            "stop_loss_pct": stop_pct,
            "target_pct": round(float(p["target"]) * 100, 2) if p["target"] is not None else None,
            "currency": p["ccy"] or "KRW",
            # 전략 보드가 '논리·반증 조건'을 보여 주는 데 쓴다(엑셀 포지션 시트 O~Q 열)
            "thesis": p.get("thesis"), "invalidation": p.get("invalidation"), "evidence": p.get("evidence"),
            "legs": [{"name": r["name"], "code": r["code"], "market": r["market"], "side": r["side"], "qty": r["qty"],
                      "entry_px": r["entry_px"], "px_now": r["px_now"], "px_date": r["px_date"],
                      "beta": r["beta"], "dur": r["dur"],
                      "mv_bn": r["now_bn"], "pnl_bn": round(r["pnl_asset_bn"] + r["pnl_fx_bn"], 2)} for r in lr],
        }
        if p["engine"] == "fx":
            notional = sum(r["now_bn"] for r in lr)
            rec.update({"base_bn": 0.0, "mv_bn": 0.0, "pnl_ytd_bn": 0.0, "hedge_pct": 0,
                        "fx_exposure_bn": round(notional, 1), "pnl_fx_ytd_bn": round(pnl_fx, 2),
                        "fx_stop_pct": stop_pct})
        else:
            rec.update({"base_bn": round(base, 1), "mv_bn": round(mv, 1), "pnl_ytd_bn": round(pnl, 2)})
            if rec["currency"] != "KRW":
                rec.update({"hedge_pct": round(float(p["hedge"] or 0) * 100), "fx_exposure_bn": round(mv, 1),
                            "pnl_fx_ytd_bn": round(pnl_fx, 2)})
            if p["engine"] == "equity":
                rec["beta"] = round(net_beta, 3)
            elif p["engine"] == "bond":
                rec.update({"duration": round(net_dur, 2), "spread_duration": float(p["sdur"] or 0),
                            "spread_beta": float(p["sbeta"] or 0)})
            elif p["engine"] == "mixed":
                eq = float(p["eq_pct"] or 0)
                dur = next((float(r["dur"]) for r in lr if r["dur"]), 0.0)
                rec["look_through"] = {"equity_pct": eq, "bond_pct": 100 - eq, "equity_beta": 1.0,
                                       "bond_duration": dur, "bond_spread_duration": float(p["sdur"] or dur),
                                       "bond_spread_beta": float(p["sbeta"] or 1.0)}
                # 전환사채 ETF 는 위탁 보고서가 없으므로 델타 가정 비율로 손익을 나눈다(추정임을 meta 에 기록)
                rec["pnl_split"] = {"equity": round(pnl * eq / 100, 2), "bond": round(pnl * (100 - eq) / 100, 2)}
        out.append(rec)
    return out, rows


def write_positions(recs: list[dict], book: Path, prices: dict) -> Path:
    as_of = datetime.date.today().isoformat()
    doc = {
        "meta": {
            "as_of": as_of, "unit": "억원",
            "source": f"{book.name} → book.py 시가 평가 ({datetime.datetime.now().isoformat(timespec='seconds')})",
            "pnl_basis": "편입일 이후 평가손익. 자산 손익(가격 변화)과 환 손익(외화 자산 미헤지분·환 오버레이)을 분리. 헤지 비용·거래비용·대차 수수료·조달비용 미반영.",
            "note": "가상 시험 포트(Active 5,000억 가상 설정 — 실제 조직·계획·보유와 무관) — 공개 상장 상품이라 종목명을 포함한다. "
                    "실제 보유로 바꿀 때는 외부 API 전송을 고려해 버킷(집계) 단위로 넣는다.",
            "price_dates": sorted({d for _, d in prices.values() if d}),
            "proxies": "KOSPI200 선물=지수 현물, 달러·엔선물=현물환, 10년 국채선물=국고10y 이론가(베이시스 무시). 전환사채(B5) 손익 분할은 델타 가정 추정.",
        },
        "positions": recs,
    }
    path = config.DATA_DIR / "positions.json"
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, path)
    return path


def write_limits(ga: dict) -> Path | None:
    lim = ga.get("_limits") or {}
    if not {"전체", "주식", "채권", "환"} <= set(lim):
        print("  경고: 엑셀 '한도' 시트를 읽지 못해 limits.json 은 건너뜀")
        return None
    base = json.loads((config.DATA_DIR / "limits.sample.json").read_text(encoding="utf-8"))
    base["meta"] = {
        "version": f"trial-{datetime.date.today().isoformat()}",
        "status": "체험용 가안 — 시험 포트(Active 5,000억, 가상) 기준. 투자위 확정 전",
        "as_of": datetime.date.today().isoformat(),
        "loss_basis": base["meta"].get("loss_basis", ""),
        "rationale": f"시험포트 엑셀 '한도' 시트: 전체 {lim['전체']:.0f}억(Active 5,000억의 {lim['전체'] / 50:.1f}%), "
                     f"주식 {lim['주식']:.0f} / 채권 {lim['채권']:.0f} / 환 {lim['환']:.0f}. 단계·손절 접근·스트레스 시나리오는 기존 가안 그대로.",
        "warnings": ["체험용 한도 — 실제 운용 개시 전 투자위원회 확정값으로 교체."],
    }
    base["total_limit_bn"] = lim["전체"]
    base["asset_limits_bn"] = {"bond": lim["채권"], "equity": lim["주식"], "fx": lim["환"]}
    base["unallocated_bn"] = lim["전체"] - lim["채권"] - lim["주식"] - lim["환"]
    base["stop_loss"]["default_pct"]["mixed"] = -3.0
    path = config.DATA_DIR / "limits.json"
    path.write_text(json.dumps(base, ensure_ascii=False, indent=1), encoding="utf-8")
    return path


def sync(book: Path = DEFAULT_BOOK, with_limits: bool = False, dry_run: bool = False, quiet: bool = False) -> dict:
    """대시보드 '일일 점검'·CLI 공통 진입점."""
    legs, pos, ga = read_book(book)
    symbols = {price_source(l) for l in legs} | {"KRW=X"}
    prices = fetch_prices(symbols)
    recs, rows = mark(legs, pos, prices, ga)
    res = {"book": book.name, "positions": len(recs), "legs": len(rows),
           "pnl_bn": round(sum(r["pnl_ytd_bn"] + r.get("pnl_fx_ytd_bn", 0) for r in recs), 2),
           "missing_prices": sorted({r["src"] for r in rows if r["px_date"] == "진입가 유지"})}
    if not quiet:
        for r in recs:
            tot = r["pnl_ytd_bn"] + r.get("pnl_fx_ytd_bn", 0)
            print(f"  {r['id']:<3} {r['short_name']:<24} 시가 {r.get('mv_bn', 0):>8.1f}억  손익 {tot:+8.2f}억  ({r['pnl_as_of']})")
        print(f"  합계 손익 {res['pnl_bn']:+.2f}억 · 가격 못 받은 종목 {res['missing_prices'] or '없음'}")
    if not dry_run:
        res["positions_path"] = str(write_positions(recs, book, prices))
        if with_limits:
            lp = write_limits(ga)
            res["limits_path"] = str(lp) if lp else None
    return res


def main():
    ap = argparse.ArgumentParser(description="보유 엑셀 → 시가 평가 → positions.json")
    ap.add_argument("book", nargs="?", default=str(DEFAULT_BOOK))
    ap.add_argument("--limits", action="store_true", help="엑셀 '한도' 시트로 data/limits.json 도 갱신(체험용)")
    ap.add_argument("--dry-run", action="store_true", help="쓰지 않고 평가만")
    a = ap.parse_args()
    res = sync(Path(a.book), with_limits=a.limits, dry_run=a.dry_run)
    print(json.dumps({k: v for k, v in res.items() if k != "missing_prices"}, ensure_ascii=False))


if __name__ == "__main__":
    if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
