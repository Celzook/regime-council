# -*- coding: utf-8 -*-
"""인포맥스 실시간 브리지 — 엑셀 애드인 함수(IMDP·IMDT)로 받은 값을 시스템으로 가져온다.

인포맥스 함수는 엑셀 안에서만 돈다. 그래서:
1. build: 브리지 엑셀(data/인포맥스_브리지.xlsx)을 만든다 — '변수' 시트 한 줄 = 지표 하나(코드·항목·실시간 값 수식),
   '캘린더' 시트 = 경제지표 발표 일정(IMDT), '검증' 시트 = 구분별 대표 항목 시험(코드·항목 이름 확인용).
2. read: 사용자가 이 파일을 열어 둘 필요는 없다. 인포맥스 프로그램(로그인)이 켜져 있으면 보이지 않는 엑셀을
   따로 띄워 값만 받아 오고 닫는다(약 10초, 사용자의 엑셀 창은 건드리지 않음). 사용자가 마침 열어 두었으면 그 값을
   읽고(읽기 전용), 둘 다 안 되면 마지막 값을 쓴다 → data/macro_live.json.
   경제지표 일정(캘린더 시트)은 인포맥스 '조회' 버튼으로 채우는 방식이라 숨김 엑셀로는 받아지지 않는다.

코드·문법 출처(이 PC 의 인포맥스 설치본에서 확인): 금리 코드·항목은 C:\\Infomax\\bin\\excel\\시가평가_matrix_3자리.xlsx
(민평 3사 BONDAVG**·"3년이하(당일)"), 환율 코드는 경제위기_비교_템플릿.xlsb(USDSP_SMBCB·DOLLAR·EURUSD·US03M·IRR:TED),
경제지표 코드는 C:\\Infomax\\bin\\code\\entry.xml(종료일 2026 이후 활성 코드만). 항목 이름은 C:\\Infomax\\bin\\ini 의
경제지표.csv(현재가_월·DATE_월·전년비_월 …)·외환_usdkrw_smbc_*.dat(현재가·중간값), 스왑포인트 코드는 excelrealmap.dat
(USD3M_SMBCB) — 10-21 숨김 엑셀로 값이 나오는 것까지 확인.

사용법:
  python infomax_bridge.py build [--force]   # 브리지 엑셀 생성(기존 파일은 --force 없이는 덮어쓰지 않음)
  python infomax_bridge.py read              # 값 읽기 → data/macro_live.json
"""
import base64
import datetime
import json
import os
import subprocess
import sys
import threading
from pathlib import Path

import config

BRIDGE = config.DATA_DIR / "인포맥스_브리지.xlsx"
LIVE = config.DATA_DIR / "macro_live.json"
SHEET = "변수"
# 경제지표 항목은 주기가 붙는다(현재가_월 …). 코드마다 주기가 달라 차례로 시도 — 값·일자 순서를 맞춘다.
ECO_ITEMS = ["현재가_월", "현재가_주", "현재가_일", "현재가_분기", "현재가_년"]
ECO_DATE_ITEMS = ["DATE_월", "DATE_주", "DATE_일", "DATE_분기", "DATE_년"]
FX_DATE_ITEMS = ["거래일", "일자"]
SCALE = {"swap_3m": 0.01, "swap_1y": 0.01}   # 스왑포인트는 1/100원 단위로 온다 → 원

# key, 이름, 담당(bond/equity/fx/all), 주제, 구분, 코드, 항목, 단위, 비고
REGISTRY = [
    # ── 국내 금리 (민평 3사, 당일) ──
    ("ktb_1y", "국고 1년", "bond", "금리", "IR", "BONDAVG01", "1년이하(당일)", "%", "민평 3사"),
    ("ktb_3y", "국고 3년", "bond", "금리", "IR", "BONDAVG01", "3년이하(당일)", "%", "민평 3사"),
    ("ktb_5y", "국고 5년", "bond", "금리", "IR", "BONDAVG01", "5년이하(당일)", "%", "민평 3사"),
    ("ktb_10y", "국고 10년", "bond", "금리", "IR", "BONDAVG01", "10년이하(당일)", "%", "민평 3사"),
    ("ktb_30y", "국고 30년", "bond", "금리", "IR", "BONDAVG01", "30년이하(당일)", "%", "민평 3사"),
    ("msb_1y", "통안 1년", "bond", "금리", "IR", "BONDAVG14", "1년이하(당일)", "%", "민평 3사 이표채"),
    ("bank_aaa_3y", "은행채 AAA 3년", "bond", "크레딧", "IR", "BONDAVG20", "3년이하(당일)", "%", "민평 3사"),
    ("card_aa_3y", "카드채 AA+ 3년", "bond", "크레딧", "IR", "BONDAVG81", "3년이하(당일)", "%", "민평 3사"),
    ("corp_aa_3y", "회사채 AA- 3년", "bond", "크레딧", "IR", "BONDAVG58", "3년이하(당일)", "%", "민평 3사 공모 무보증"),
    ("corp_a_3y", "회사채 A+ 3년", "bond", "크레딧", "IR", "BONDAVG59", "3년이하(당일)", "%", "민평 3사 공모 무보증"),
    ("corp_bbb_3y", "회사채 BBB- 3년", "bond", "크레딧", "IR", "BONDAVG64", "3년이하(당일)", "%", "민평 3사 공모 무보증"),
    ("ust_3m", "미국채 3개월", "bond", "금리", "IR", "US03M", "현재가", "%", "템플릿 확인 코드"),
    ("ted", "TED 스프레드", "bond", "신용·유동성", "FRN", "IRR:TED", "현재가", "bp", "2025-02 이후 갱신 없음 — 참고만"),
    # ── 환율 (실시간) ──
    ("usdkrw", "USD/KRW 현물", "fx", "환율", "FX", "USDSP_SMBCB_EXT", "현재가", "원", "서울외환 은행간 · 24시간"),
    ("dxy", "달러인덱스", "fx", "환율", "FX", "DOLLAR", "현재가", "pt", "실시간"),
    ("eurusd", "EUR/USD", "fx", "환율", "FX", "EURUSD", "현재가", "달러", "실시간"),
    ("swap_3m", "원/달러 스왑포인트 3M", "fx", "헤지비용", "FX", "USD3M_SMBCB", "중간값", "원", "서울외환 중간값(1/100원 → 원 환산)"),
    ("swap_1y", "원/달러 스왑포인트 1Y", "fx", "헤지비용", "FX", "USD1Y_SMBCB", "중간값", "원", "서울외환 중간값(1/100원 → 원 환산)"),
    # ── 국내 거시 (경제지표) ──
    ("kr_cpi_yoy", "소비자물가 전년비", "bond", "물가", "ECO", "701978", "", "%", "월"),
    ("kr_unemp", "실업률", "bond", "고용", "ECO", "27", "", "%", "월"),
    ("kr_employed", "취업자", "equity", "고용", "ECO", "30", "", "천명", "월"),
    ("kr_leading", "선행지수 순환변동치", "equity", "경기", "ECO", "702165", "", "p", "월"),
    ("kr_exports", "수출액", "equity", "경기", "ECO", "557150", "", "천USD", "월"),
    ("kr_exports_20d", "수출액(20일)", "equity", "경기", "ECO", "705150", "", "천USD", "월 중순 속보"),
    ("kr_ca", "경상수지", "fx", "대외", "ECO", "10", "", "백만USD", "월"),
    ("kr_fxres", "외환보유액", "fx", "대외", "ECO", "400020", "", "천USD", "월"),
    ("kr_csi", "소비자심리지수", "equity", "심리", "ECO", "701542", "", "pt", "월"),
    ("kr_hh_loans", "가계대출(예금취급기관)", "bond", "신용·유동성", "ECO", "703216", "", "십억원", "월"),
    # ── 해외 거시 ──
    ("us_cpi", "미국 CPI 전년비", "bond", "물가", "ECO", "135012", "전년비_월", "%", "월 · 계절조정 지수의 전년비"),
    ("us_core_cpi", "미국 근원 CPI 전년비", "bond", "물가", "ECO", "135013", "전년비_월", "%", "월 · 계절조정 지수의 전년비"),
    ("us_pce_yoy", "미국 PCE 전년비", "bond", "물가", "ECO", "705121", "", "%", "월"),
    ("us_nfp", "미국 비농업 고용자수", "bond", "고용", "ECO", "130015", "", "천명", "월"),
    ("us_unemp", "미국 실업률", "bond", "고용", "ECO", "130001", "", "%", "월"),
    ("us_claims", "미국 신규 실업수당청구", "bond", "고용", "ECO", "130002", "", "천명", "주"),
    ("us_ism_mfg", "미국 ISM 제조업", "equity", "경기", "ECO", "131001", "", "pt", "월"),
    ("us_ism_svc", "미국 ISM 비제조업", "equity", "경기", "ECO", "131002", "", "pt", "월"),
    ("us_retail", "미국 소매매출", "equity", "소비", "ECO", "132020", "", "백만USD", "월"),
    ("cn_pmi", "중국 제조업 PMI", "equity", "경기", "ECO", "702440", "", "pt", "월"),
    ("ez_cpi", "유로존 소비자물가 전년비", "bond", "물가", "ECO", "700344", "전년비_월", "%", "월 · 지수의 전년비"),
]
COLS = ["키", "이름", "담당", "주제", "구분", "코드", "항목", "단위", "비고", "값", "일자", "상태"]


def _value_formula(r: int) -> str:
    """항목이 적혀 있으면 그 항목으로 IMDP, ECO 인데 비어 있으면 주기 후보를 차례로 시도. 코드가 비면 빈칸."""
    eco = ",".join(f'IFERROR(1*_xll.IMDP("ECO",$F{r},"{it}")' for it in ECO_ITEMS) + ',""' + ")" * len(ECO_ITEMS)
    plain = f'IFERROR(1*_xll.IMDP($E{r},$F{r},$G{r}),"")'
    return f'=IF($F{r}="","",IF(AND($E{r}="ECO",$G{r}=""),{eco},{plain}))'


def _date_formula(r: int) -> str:
    def chain(g, items):  # 숫자로 나오는 첫 항목(… ·빈칸은 건너뜀)
        return ",".join(f'IFERROR(1*_xll.IMDP({g},$F{r},"{it}")' for it in items) + ',""' + ")" * len(items)
    return (f'=IF($F{r}="","",IF($E{r}="ECO",{chain(chr(34) + "ECO" + chr(34), ECO_DATE_ITEMS)},'
            f'IF($E{r}="IR",IFERROR(1*_xll.IMDP($E{r},$F{r},"날짜(당일)"),""),{chain(f"$E{r}", FX_DATE_ITEMS)})))')


def build(force: bool = False) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    if BRIDGE.exists() and not force:
        raise SystemExit(f"{BRIDGE.name} 이 이미 있습니다(코드를 고쳤을 수 있음). 다시 만들려면 --force")
    F = "Arial"
    head = Font(name=F, bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="3A3A3C")
    inp = Font(name=F, color="0000FF"); base = Font(name=F)
    wb = Workbook()
    ws = wb.active; ws.title = SHEET
    ws["A1"] = "인포맥스 브리지 — 열어 둘 필요 없습니다. 인포맥스가 켜져 있으면 시스템이 보이지 않는 엑셀로 이 파일을 열어 값을 받아 갑니다."
    ws["A1"].font = Font(name=F, bold=True)
    ws["A2"] = "파란 글씨(구분·코드·항목)는 고칠 수 있습니다. 코드는 인포맥스 엑셀 애드인의 코드 검색으로 찾아 넣으세요. 행을 추가해도 됩니다(값·일자 수식은 위 행을 복사)."
    ws["A2"].font = Font(name=F, size=9, color="666666")
    for c, h in enumerate(COLS, 1):
        cell = ws.cell(4, c, h); cell.font, cell.fill = head, fill
    for i, row in enumerate(REGISTRY, start=5):
        for c, v in enumerate(row, 1):
            cell = ws.cell(i, c, v); cell.font = inp if c in (5, 6, 7) else base
        ws.cell(i, 10, _value_formula(i)).font = base
        ws.cell(i, 11, _date_formula(i)).font = base
        ws.cell(i, 12, f'=IF($F{i}="","코드 입력 필요",IF(ISNUMBER($J{i}),"OK","값 없음 — 코드·항목 확인"))').font = base
    for col, w in zip("ABCDEFGHIJKL", (14, 24, 8, 11, 6, 14, 15, 9, 30, 12, 12, 22)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"

    cal = wb.create_sheet("캘린더")
    items = ["발표일시", "기준일자", "내용", "실제치", "주기", "국가", "발표처", "단위", "종목코드"]
    cal["A1"] = "경제지표 발표 일정 (IMDT) — 최근 3일 ~ 향후 14일, 한국·미국·중국·유럽·일본"; cal["A1"].font = Font(name=F, bold=True)
    for c, h in enumerate(items, 1):
        cal.cell(2, c, h).font = head; cal.cell(2, c).fill = fill
    cal["K2"] = "조회 기간"; cal["L2"] = '=TEXT(TODAY()-3,"yyyymmdd")&"~"&TEXT(TODAY()+14,"yyyymmdd")'
    cal["K3"] = "국가"; cal["L3"] = "한국|미국|중국|유럽|일본"; cal["L3"].font = inp
    # 결과는 수식 셀(A3)에서 2행 아래(A5)부터 펼쳐진다(Pos=20)
    cal["A3"] = '=_xll.IMDT("경제지표","CALENDAR",$A$2:$I$2,80,$L$2,$L$3,"발표일시","Sort=A,DtFmt=1,TmFmt=1,Pos=20")'
    for col, w in zip("ABCDEFGHI", (17, 11, 34, 10, 6, 8, 16, 9, 10)):
        cal.column_dimensions[col].width = w

    chk = wb.create_sheet("검증")
    chk["A1"] = "항목 이름 확인용 — 숫자가 나오는 칸이 맞는 항목입니다(경제지표는 문서에 예시가 없어 후보를 시험)."; chk["A1"].font = Font(name=F, bold=True)
    for c, h in enumerate(["시험", "구분", "코드", "항목", "결과"], 1):
        chk.cell(3, c, h).font = head; chk.cell(3, c).fill = fill
    tests = [("ECO 값", "ECO", "701978", "현재가_월"), ("ECO 일자", "ECO", "701978", "DATE_월"), ("ECO 전년비", "ECO", "135012", "전년비_월"),
             ("IR 값", "IR", "BONDAVG01", "3년이하(당일)"), ("IR 일자", "IR", "BONDAVG01", "날짜(당일)"),
             ("FX 값", "FX", "USDSP_SMBCB_EXT", "현재가"), ("FX 스왑", "FX", "USD3M_SMBCB", "중간값"),
             ("FX 일자", "FX", "USD3M_SMBCB", "거래일"), ("FRN 값", "FRN", "IRR:TED", "현재가")]
    for i, (t, g, code, it) in enumerate(tests, start=4):
        for c, v in enumerate((t, g, code, it), 1):
            chk.cell(i, c, v).font = base
        chk.cell(i, 5, f'=IFERROR(_xll.IMDP(B{i},C{i},D{i}),"오류")').font = base
    for col, w in zip("ABCDE", (10, 6, 14, 16, 20)):
        chk.column_dimensions[col].width = w
    for s in wb.worksheets:
        s.sheet_view.showGridLines = False
        for row in s.iter_rows(min_row=5):
            for cell in row:
                cell.alignment = Alignment(vertical="center")
    BRIDGE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(BRIDGE)
    return BRIDGE


# ── 읽기 ────────────────────────────────────────────────────────
_PS = r"""
[Console]::OutputEncoding = [System.Text.Encoding]::UTF8
try { $xl = [Runtime.InteropServices.Marshal]::GetActiveObject('Excel.Application') } catch { '{"open":false,"reason":"excel"}'; exit 0 }
$wb = $null
foreach ($b in $xl.Workbooks) { if ($b.Name -like '__NAME__*') { $wb = $b } }
if (-not $wb) { '{"open":false,"reason":"book"}'; exit 0 }
$res = [ordered]@{ open = $true; path = $wb.FullName; sheets = [ordered]@{} }
foreach ($n in @('__S1__', '__S2__', '__S3__')) {
  try {
    $v = $wb.Worksheets.Item($n).UsedRange.Value2
    $rows = New-Object System.Collections.ArrayList
    if ($v -is [System.Array]) {
      for ($i = 1; $i -le $v.GetLength(0); $i++) {
        $r = New-Object System.Collections.ArrayList
        for ($j = 1; $j -le $v.GetLength(1); $j++) { [void]$r.Add($v.GetValue($i, $j)) }
        [void]$rows.Add($r)
      }
    }
    $res.sheets[$n] = $rows
  } catch { $res.sheets[$n] = $null }
}
$res | ConvertTo-Json -Depth 6 -Compress
"""


def _read_live() -> dict | None:
    """실행 중인 엑셀에서 브리지 통합문서 값만 읽는다(쓰기·재계산 없음). 열려 있지 않으면 None."""
    script = (_PS.replace("__NAME__", BRIDGE.stem).replace("__S1__", SHEET)
              .replace("__S2__", "캘린더").replace("__S3__", "검증"))
    enc = base64.b64encode(script.encode("utf-16-le")).decode()
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-EncodedCommand", enc],
                           capture_output=True, timeout=25)
    except (subprocess.TimeoutExpired, OSError):
        return None
    out = r.stdout.decode("utf-8", errors="replace").strip()
    try:
        d = json.loads(out[out.find("{"):]) if "{" in out else None
    except json.JSONDecodeError:
        return None
    return d if d and d.get("open") else None


def _read_saved() -> dict | None:
    """마지막으로 저장된 값(엑셀이 저장할 때 남긴 계산 결과)."""
    if not BRIDGE.exists():
        return None
    from openpyxl import load_workbook
    wb = load_workbook(BRIDGE, data_only=True, read_only=True)
    sheets = {}
    for n in (SHEET, "캘린더", "검증"):
        if n in wb.sheetnames:
            sheets[n] = [list(r) for r in wb[n].iter_rows(values_only=True)]
    return {"open": False, "path": str(BRIDGE), "sheets": sheets}


# ── 백그라운드(숨김 엑셀) 읽기 ──────────────────────────────────────
# 사용자가 브리지를 열어 두지 않아도 되게: 보이지 않는 엑셀을 따로 하나 띄워 애드인을 올리고, 브리지를 읽기 전용으로
# 열어 값이 들어오면 읽고 닫는다. 사용자의 엑셀(이미 떠 있는 프로세스)에는 붙지 않는다 — 새로 뜬 프로세스가
# 기존 PID 와 같으면 아무것도 하지 않고 빠진다. 애드인은 C:\Infomax\bin\Excel64 를 작업 폴더로 해야 의존 DLL 을 찾는다.
XLL = Path(r"C:\Infomax\bin\Excel64\IMxlExcelAi64.xll")
PIDFILE = config.DATA_DIR / ".infomax_excel.pid"
LOCK = threading.Lock()  # 엑셀 COM 호출은 한 번에 하나(웹 화면·일일 점검이 함께 쓴다)

_PS_HEADLESS = r"""
[Console]::OutputEncoding = [System.Text.Encoding]::UTF8
$ProgressPreference = 'SilentlyContinue'
$res = [ordered]@{ open = $false; reason = ''; path = '__PATH__'; sheets = [ordered]@{}; xpid = 0; secs = 0; note = '' }
$t0 = Get-Date
if (-not (Get-Process -Name infomaxmain -ErrorAction SilentlyContinue)) { $res.reason = 'infomax' }
else {
  $before = @(Get-Process -Name EXCEL -ErrorAction SilentlyContinue | ForEach-Object { $_.Id })
  Add-Type -Namespace RC -Name Win -MemberDefinition '[DllImport("user32.dll")] public static extern uint GetWindowThreadProcessId(IntPtr h, out uint p);'
  $xl = $null; $wb = $null; $ws = $null; $quit = $true
  try {
    $xl = New-Object -ComObject Excel.Application
    [uint32]$p = 0; [void][RC.Win]::GetWindowThreadProcessId([IntPtr]$xl.Hwnd, [ref]$p)
    if ($before -contains [int]$p) { $quit = $false; $xl = $null; throw 'attached' }
    $res.xpid = [int]$p; Set-Content -Path '__PIDFILE__' -Value ("{0} {1}" -f $p, $t0.ToString('o')) -Encoding ascii
    $xl.Visible = $false; $xl.DisplayAlerts = $false; $xl.ScreenUpdating = $false
    $xl.IgnoreRemoteRequests = $true; $xl.AutomationSecurity = 3
    [void]$xl.ExecuteExcel4Macro('DIRECTORY("__XLLDIR__")')
    if (-not $xl.RegisterXLL('__XLL__')) { throw 'xll' }
    $m = [Type]::Missing
    $wb = $xl.Workbooks.Open('__PATH__', 0, $true, $m, $m, $m, $true, $m, $m, $false, $false, $m, $false)
    $xl.CalculateFull()
    $ws = $wb.Worksheets.Item('__WATCH__')
    $last = -1; $stable = 0; $deadline = $t0.AddSeconds(__WAIT__)
    while ((Get-Date) -lt $deadline) {
      Start-Sleep -Milliseconds 1000
      try { $v = $ws.UsedRange.Value2 } catch { continue }
      $c = 0; foreach ($x in $v) { if ($x -is [double]) { $c++ } }
      if ($c -eq $last) { $stable++ } else { $stable = 0; $last = $c }
      if ($c -gt 0 -and $stable -ge __STABLE__) { break }
    }
    foreach ($n in @(__SHEETS__)) {
      try {
        $v = $wb.Worksheets.Item($n).UsedRange.Value2
        $rows = New-Object System.Collections.ArrayList
        if ($v -is [System.Array]) {
          for ($i = 1; $i -le $v.GetLength(0); $i++) {
            $r = New-Object System.Collections.ArrayList
            for ($j = 1; $j -le $v.GetLength(1); $j++) { [void]$r.Add($v.GetValue($i, $j)) }
            [void]$rows.Add($r)
          }
        }
        $res.sheets[$n] = $rows
      } catch { $res.sheets[$n] = $null }
    }
    $res.open = $true
  } catch { $res.reason = $_.Exception.Message }
  finally {
    try { if ($xl) { $xl.IgnoreRemoteRequests = $false } } catch {}
    try { if ($wb) { $wb.Close($false) } } catch {}
    try { if ($xl -and $xl.Workbooks.Count -gt 0) { $quit = $false; $xl.ScreenUpdating = $true; $xl.Visible = $true; $res.note = 'foreign' } } catch {}
    try { if ($xl -and $quit) { $xl.Quit() } } catch {}
    foreach ($o in @($ws, $wb, $xl)) { if ($o) { try { [void][Runtime.InteropServices.Marshal]::ReleaseComObject($o) } catch {} } }
    $ws = $null; $wb = $null; $xl = $null; [GC]::Collect(); [GC]::WaitForPendingFinalizers()
  }
}
$res.secs = [math]::Round(((Get-Date) - $t0).TotalSeconds, 1)
$res | ConvertTo-Json -Depth 6 -Compress
"""

# 숨김 엑셀이 스스로 안 닫히면 정리: 우리가 띄운 PID(시작 시각까지 확인)와 그 자식 애드인 통신 프로세스만 종료.
_PS_REAP = r"""
$ProgressPreference = 'SilentlyContinue'
$xpid = __PID__; $since = [datetime]::Parse('__SINCE__').AddSeconds(-5)
for ($k = 0; $k -lt 30 -and (Get-Process -Id $xpid -ErrorAction SilentlyContinue); $k++) { Start-Sleep -Milliseconds 500 }
$p = Get-Process -Id $xpid -ErrorAction SilentlyContinue
if ($p -and $p.ProcessName -eq 'EXCEL' -and $p.StartTime -ge $since) {
  $kids = @(Get-CimInstance Win32_Process -Filter "ParentProcessId=$xpid" | Where-Object { $_.Name -like 'IMxl*' } | ForEach-Object { $_.ProcessId })
  Stop-Process -Id $xpid -Force -ErrorAction SilentlyContinue
  Start-Sleep -Seconds 2
  foreach ($k in $kids) {
    $g = @(Get-CimInstance Win32_Process -Filter "ParentProcessId=$k" | Where-Object { $_.Name -like 'IMxl*' } | ForEach-Object { $_.ProcessId })
    foreach ($q in @($k) + $g) { Stop-Process -Id $q -Force -ErrorAction SilentlyContinue }
  }
  'killed'
} else { 'exited' }
"""


def _ps(script: str, timeout: float) -> str | None:
    enc = base64.b64encode(script.encode("utf-16-le")).decode()
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-EncodedCommand", enc],
                           capture_output=True, timeout=timeout)
    except (subprocess.TimeoutExpired, OSError):
        return None
    return r.stdout.decode("utf-8", errors="replace").strip()


def _reap() -> str:
    """숨김 엑셀 정리(정상 종료를 15초 기다린 뒤에도 남아 있으면 그 PID 만 종료)."""
    try:
        pid, since = PIDFILE.read_text(encoding="ascii").split()
        int(pid)
    except (OSError, ValueError):
        return "none"
    out = _ps(_PS_REAP.replace("__PID__", pid).replace("__SINCE__", since), timeout=40) or ""
    PIDFILE.unlink(missing_ok=True)
    return out.splitlines()[-1] if out else "unknown"


def infomax_running() -> bool:
    out = _ps("if (Get-Process -Name infomaxmain -ErrorAction SilentlyContinue) { 'y' } else { 'n' }", timeout=15)
    return out == "y"


def read_headless(path: Path = BRIDGE, sheets=(SHEET, "검증"), watch: str = SHEET,
                  wait: int = 40, stable: int = 4) -> dict | None:
    """숨김 엑셀로 통합문서를 열어 애드인 값이 들어오면 읽는다. 인포맥스 미실행·실패 시 reason 이 담긴 dict."""
    if not XLL.exists():
        return {"open": False, "reason": "xll-missing"}
    script = (_PS_HEADLESS.replace("__PATH__", str(path)).replace("__PIDFILE__", str(PIDFILE))
              .replace("__XLLDIR__", str(XLL.parent)).replace("__XLL__", str(XLL)).replace("__WATCH__", watch)
              .replace("__WAIT__", str(wait)).replace("__STABLE__", str(stable))
              .replace("__SHEETS__", ", ".join(f"'{s}'" for s in sheets)))
    try:
        out = _ps(script, timeout=wait + 45)
    finally:
        reap = _reap()
    if not out or "{" not in out:
        return {"open": False, "reason": "timeout", "reap": reap}
    try:
        d = json.loads(out[out.find("{"):])
    except json.JSONDecodeError:
        return {"open": False, "reason": "parse", "reap": reap}
    d["reap"] = reap
    return d


def _xl_date(v):
    """엑셀 일련번호·YYYYMMDD·문자열 날짜를 ISO 로."""
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        if 19000101 <= v <= 21001231:
            s = str(int(v)); return f"{s[:4]}-{s[4:6]}-{s[6:]}"
        if 20000 < v < 80000:
            return (datetime.date(1899, 12, 30) + datetime.timedelta(days=float(v))).isoformat()
        return None
    s = str(v).strip().replace(".", "-").replace("/", "-")
    if len(s) == 8 and s.isdigit():
        return f"{s[:4]}-{s[4:6]}-{s[6:]}"
    return s[:10] or None


def parse(raw: dict) -> dict:
    sh = raw.get("sheets") or {}
    rows = sh.get(SHEET) or []
    hdr = next((i for i, r in enumerate(rows) if r and r[0] == "키"), None)
    series = []
    if hdr is not None:
        for r in rows[hdr + 1:]:
            r = list(r) + [None] * (len(COLS) - len(r))
            if not r[0]:
                continue
            val = r[9] if isinstance(r[9], (int, float)) and not isinstance(r[9], bool) else None
            if val is not None:
                val = round(val * SCALE.get(r[0], 1), 4)
            series.append({"key": r[0], "name": r[1], "desk": r[2], "theme": r[3], "gubun": r[4], "code": r[5],
                           "item": r[6], "unit": r[7], "note": r[8], "value": val, "date": _xl_date(r[10]),
                           "status": "OK" if val is not None else ("코드 입력 필요" if not r[5] else "값 없음 — 코드·항목 확인")})
    cal = []
    crow = sh.get("캘린더") or []
    for r in crow[4:]:          # 5행부터 결과(Pos=20)
        if not r or not any(x not in (None, "") for x in r[:9]):
            continue
        r = list(r) + [None] * 9
        cal.append({"when": r[0], "ref": _xl_date(r[1]) or r[1], "title": r[2], "actual": r[3], "freq": r[4],
                    "country": r[5], "source": r[6], "unit": r[7], "code": r[8]})
    return {"series": series, "calendar": cal[:80]}


def read(headless: bool = True) -> dict:
    """열린 브리지 → (인포맥스가 켜져 있으면) 숨김 엑셀 → 마지막 값 순서. 다른 읽기가 진행 중이면 RuntimeError."""
    if not LOCK.acquire(blocking=False):
        raise RuntimeError("이미 읽는 중입니다.")
    try:
        return _read_locked(headless)
    finally:
        LOCK.release()


def _read_locked(headless: bool) -> dict:
    now = datetime.datetime.now().isoformat(timespec="seconds")
    raw, mode, why = _read_live(), "실시간(열린 엑셀)", ""
    if raw is None and headless and BRIDGE.exists():
        h = read_headless()
        if h and h.get("open"):
            raw, mode = h, f"자동(숨김 엑셀 {h.get('secs')}초)"
        else:
            why = {"infomax": "인포맥스 프로그램이 꺼져 있음", "xll-missing": "인포맥스 엑셀 애드인을 찾지 못함",
                   "attached": "새 엑셀을 띄우지 못함"}.get((h or {}).get("reason"), f"자동 읽기 실패({(h or {}).get('reason')})")
    if raw is None:
        prev = load()
        if prev.get("series"):                      # 마지막으로 받은 값을 유지(받은 시각은 그대로)
            prev["meta"]["status"] = why or "마지막 값"
            prev["meta"]["checked_at"] = now
            _write(prev)
            return prev
        raw = _read_saved(); mode = "저장값(파일)"
    if raw is None:
        doc = {"meta": {"read_at": now, "checked_at": now, "mode": "없음", "status": why,
                        "note": f"{BRIDGE.name} 이 없습니다 — python infomax_bridge.py build"}, "series": [], "calendar": []}
    else:
        p = parse(raw)
        ok = sum(1 for s in p["series"] if s["status"] == "OK")
        doc = {"meta": {"read_at": now, "checked_at": now, "mode": mode, "status": why or "정상", "workbook": raw.get("path"),
                        "ok": ok, "total": len(p["series"]),
                        "note": "인포맥스 데이터 — 사내 화면·코드 계산용. 외부 AI 전송은 라이선스·준법 확인 후(config.SHARE_VENDOR_DATA)."},
               "series": p["series"], "calendar": p["calendar"] or load().get("calendar", [])}
    _write(doc)
    return doc


def _write(doc: dict) -> None:
    tmp = LIVE.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(doc, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    os.replace(tmp, LIVE)


def age_minutes() -> float | None:
    """마지막으로 값을 확인한 뒤 지난 분(없으면 None)."""
    m = load().get("meta", {})
    try:
        return (datetime.datetime.now() - datetime.datetime.fromisoformat(m.get("checked_at") or m.get("read_at"))).total_seconds() / 60
    except (TypeError, ValueError):
        return None


def refresh_if_stale(max_age_min: float) -> bool:
    """값이 오래됐으면 백그라운드에서 다시 읽는다. 지금 읽는 중이면 True."""
    if LOCK.locked():
        return True
    a = age_minutes()
    if a is not None and a < max_age_min:
        return False

    def _run():
        try:
            read()
        except Exception as e:  # noqa: BLE001 — 화면 갱신 실패가 서버를 멈추지 않게
            print(f"  인포맥스 자동 읽기 실패: {e}")
    threading.Thread(target=_run, daemon=True).start()
    return True


def load() -> dict:
    try:
        return json.loads(LIVE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {"meta": {"mode": "없음"}, "series": [], "calendar": []}


if __name__ == "__main__":
    if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    cmd = sys.argv[1] if len(sys.argv) > 1 else "read"
    if cmd == "build":
        print("생성:", build(force="--force" in sys.argv))
    else:
        d = read(); m = d["meta"]
        print(f"{m.get('mode')} · {m.get('status')} · OK {m.get('ok', 0)}/{m.get('total', 0)} · 캘린더 {len(d['calendar'])}건")
        for s in d["series"][:50]:
            print(f"  {s['status']:<16} {s['name']:<22} {s['value']} {s['unit'] or ''} {s['date'] or ''}")
